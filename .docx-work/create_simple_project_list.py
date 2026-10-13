import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


BASE_DIR = Path(r"C:\UR\Instrukcje\Projektowanie Aplikacji Webowych Opartych o Mikrousługi\.docx-work")
BUILDER_PATH = BASE_DIR / "create_project_idea_form.py"

spec = importlib.util.spec_from_file_location("project_form_builder", BUILDER_PATH)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


def add_table(doc):
    headers = ["Imię i nazwisko", "Nr albumu", "Pomysł na projekt", "Problem / potrzeba", "Odbiorca / aktor", "Pojęcia domenowe"]
    placeholder = ["Wpisz imię i nazwisko", "Wpisz numer", "Wpisz temat", "Jaki problem rozwiązujecie?", "Kto będzie korzystał?", "Wypisz 3–5 pojęć"]
    rows = 5
    table = doc.add_table(rows=1 + rows, cols=6)
    builder.prepare_table(table, [2.0, 1.2, 2.0, 2.1, 1.5, 1.39])

    for cell, text in zip(table.rows[0].cells, headers):
        builder.set_cell_fill(cell, builder.BLUE)
        builder.set_cell_content(cell, text, "UR Nagłówek tabeli", color="FFFFFF", bold=True, size=9)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.rows[0].height = Inches(0.55)
        table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    example = [
        "PRZYKŁAD",
        "—",
        "Rezerwacja sal",
        "Student chce znaleźć i zarezerwować wolną salę.",
        "student, prowadzący",
        "sala, rezerwacja, termin, użytkownik",
    ]
    for row_index, row in enumerate(table.rows[1:]):
        builder.set_row_cant_split(row)
        row.height = Inches(0.60)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        values = example if row_index == 0 else placeholder
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            builder.set_cell_fill(cell, builder.RESPONSE_FILL)
            if row_index == 0:
                builder.set_cell_content(cell, text, "UR Table Body", bold=(index == 0), size=8.5)
            else:
                builder.set_cell_content(cell, text, "UR Placeholder", color=builder.GRAY, italic=True, size=9)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def build():
    doc = Document(str(builder.REFERENCE))
    builder.remove_body_content(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    doc.core_properties.title = "Pomysły na projekty"
    doc.core_properties.subject = "LAB1 — Projektowanie aplikacji webowych opartych o mikrousługi"
    doc.core_properties.author = ""

    title = doc.add_paragraph(style="Title")
    builder.fill_paragraph(title, "Pomysły na projekty", "Title", color=builder.BLUE, bold=True, size=22)
    subtitle = doc.add_paragraph(style="UR Subtitle")
    builder.fill_paragraph(subtitle, "Projektowanie aplikacji webowych opartych o mikrousługi  |  LAB1", "UR Subtitle", color=builder.GRAY, bold=True, size=10.5)
    instruction = doc.add_paragraph(style="UR Instruction")
    builder.fill_paragraph(instruction, "Zacznijcie od domeny, nie od technologii. Wpiszcie problem, odbiorcę lub aktora oraz 3–5 pojęć domenowych. Pojęcia domenowe to nazwy osób, rzeczy, zdarzeń lub stanów, np. Sala, Rezerwacja, Termin i Użytkownik.", "UR Instruction", color=builder.GRAY, size=9.5)
    builder.add_spacer(doc)
    add_table(doc)

    footer = doc.sections[0].footer
    if footer.paragraphs:
        builder.fill_paragraph(footer.paragraphs[0], "Projektowanie aplikacji webowych opartych o mikrousługi  |  LAB1  |  Pomysły na projekty", "Footer", color=builder.BLUE, size=8)

    builder.OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(builder.OUTPUT))


if __name__ == "__main__":
    build()
