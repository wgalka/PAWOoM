from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(r"C:\UR\Instrukcje\ProjektInżynierskiSZwI\2026-2027\lab1\Karta projektu.docx")
OUTPUT = Path(r"C:\UR\Instrukcje\Projektowanie Aplikacji Webowych Opartych o Mikrousługi\2026_2027\karta-pomyslu-projektu.docx")

BLUE = "0033A0"
LABEL_FILL = "EAF2FB"
RESPONSE_FILL = "F6F9FD"
BORDER = "D9D9D9"
GRAY = "606A73"


def remove_body_content(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=150, bottom=100, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_row_cant_split(row, value=True):
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if value and node is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    elif not value and node is not None:
        tr_pr.remove(node)


def clear_cell(cell):
    cell._tc.clear_content()
    cell.add_paragraph()


def add_run_with_font(run, name="Corbel"):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def fill_paragraph(paragraph, text, style, color=None, bold=None, italic=None, size=None):
    paragraph.style = style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    add_run_with_font(run)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    return paragraph


def prepare_table(table, widths):
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for row in table.rows:
        set_row_cant_split(row)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.width = Inches(widths[index])


def set_cell_content(cell, text, style="UR Wpis studenta", color=None, bold=None, italic=None, size=None):
    clear_cell(cell)
    p = cell.paragraphs[0]
    fill_paragraph(p, text, style, color=color, bold=bold, italic=italic, size=size)
    return p


def add_prompt_lines(cell, prompt, lines=2):
    clear_cell(cell)
    p = cell.paragraphs[0]
    fill_paragraph(p, prompt, "UR Placeholder", color=GRAY, italic=True)
    for _ in range(lines):
        p = cell.add_paragraph()
        fill_paragraph(p, "\u00a0", "UR Wpis studenta")


def add_two_col_table(doc, label, prompt, lines=2):
    table = doc.add_table(rows=1, cols=2)
    prepare_table(table, [1.7319, 5.0396])
    set_cell_fill(table.cell(0, 0), LABEL_FILL)
    set_cell_fill(table.cell(0, 1), RESPONSE_FILL)
    set_cell_content(table.cell(0, 0), label, "UR Etykieta")
    add_prompt_lines(table.cell(0, 1), prompt, lines)
    return table


def add_three_col_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=3)
    prepare_table(table, [2.2576, 2.2576, 2.2576])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_fill(cell, BLUE)
        set_cell_content(cell, header, "UR Nagłówek tabeli", color="FFFFFF", bold=True, size=9)
    for row, values in zip(table.rows[1:], rows):
        set_row_cant_split(row)
        for cell, value in zip(row.cells, values):
            set_cell_fill(cell, RESPONSE_FILL)
            set_cell_content(cell, value, "UR Table Body", size=9)
    return table


def add_spacer(doc):
    p = doc.add_paragraph(style="UR Spacer")
    p.add_run(" ")


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    fill_paragraph(p, text, "Heading 1", color=BLUE, bold=True, size=12.5)
    return p


def add_instruction(doc, text):
    p = doc.add_paragraph(style="UR Instruction")
    fill_paragraph(p, text, "UR Instruction", color=GRAY, size=9.5)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break()


def build():
    doc = Document(str(REFERENCE))
    remove_body_content(doc)

    props = doc.core_properties
    props.title = "Karta pomysłu na projekt"
    props.subject = "LAB1 — Projektowanie aplikacji webowych opartych o mikrousługi"
    props.author = ""

    title = doc.add_paragraph(style="Title")
    fill_paragraph(title, "Karta pomysłu na projekt", "Title", color=BLUE, bold=True, size=22)
    subtitle = doc.add_paragraph(style="UR Subtitle")
    fill_paragraph(subtitle, "Projektowanie aplikacji webowych opartych o mikrousługi  |  LAB1", "UR Subtitle", color=GRAY, bold=True, size=10.5)
    add_instruction(doc, "Jak wypełnić: wybierz konkretny problem, opisz jego odbiorców i zaproponuj pierwszą usługę. Nie musisz jeszcze znać wszystkich technologii. Ważniejsze jest jasne rozumowanie niż gotowa nazwa aplikacji.")

    add_heading(doc, "Metryka pomysłu")
    add_two_col_table(doc, "Zespół lub grupa", "Wpisz numer grupy, nazwę zespołu albo oznaczenie osoby / pary.", lines=1)
    add_spacer(doc)
    add_two_col_table(doc, "Data", "Wpisz datę opracowania pomysłu.", lines=1)
    add_spacer(doc)
    add_two_col_table(doc, "Nazwa robocza", "Wpisz krótką nazwę, która wskazuje temat projektu. Nazwa może się później zmienić.", lines=1)

    add_heading(doc, "1  Problem i odbiorca")
    add_instruction(doc, "Opisz sytuację, którą chcesz poprawić. Unikaj ogólnych haseł typu „aplikacja do wszystkiego”.")
    add_spacer(doc)
    add_two_col_table(doc, "Problem", "Jaka konkretna trudność występuje dzisiaj? Kogo dotyczy i jakie ma konsekwencje?", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Odbiorca rozwiązania", "Kto będzie korzystał z systemu albo jego wyników? Opisz jedną główną grupę odbiorców.", lines=3)
    add_spacer(doc)
    add_two_col_table(doc, "Najważniejszy scenariusz", "Opisz jedną sytuację: kto, co chce zrobić, z jakim rezultatem.", lines=3)
    add_spacer(doc)
    add_two_col_table(doc, "Wartość rozwiązania", "Co poprawi się dzięki systemowi? Jaką czynność, decyzję lub proces będzie można wykonać lepiej?", lines=3)

    add_page_break(doc)
    add_heading(doc, "2  Domena i granice odpowiedzialności")
    add_instruction(doc, "Najpierw nazwij pojęcia domenowe, a dopiero potem zastanów się, które z nich powinny należeć do jednej usługi.")
    add_spacer(doc)
    add_two_col_table(doc, "Pojęcia domenowe", "Wypisz 3–5 pojęć używanych w problemie i krótko wyjaśnij ich znaczenie.", lines=5)
    add_spacer(doc)
    add_two_col_table(doc, "Proponowana pierwsza usługa", "Podaj nazwę usługi i jednym zdaniem opisz jej odpowiedzialność biznesową.", lines=3)
    add_spacer(doc)
    add_two_col_table(doc, "Usługa odpowiada za", "Wypisz dane, operacje i reguły, które powinny należeć do pierwszej usługi.", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Usługa nie odpowiada za", "Wypisz elementy systemu, które świadomie pozostawiasz poza pierwszą usługą.", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Dlaczego nie tylko moduł", "Czy pierwsza usługa mogłaby działać jako osobny proces i komunikować się przez API? Podaj argument.", lines=3)

    add_heading(doc, "3  Plan usług i własność danych")
    add_instruction(doc, "Zaplanuj co najmniej dwie usługi przyszłego systemu. Na LAB1 implementujesz tylko pierwszą.")
    add_spacer(doc)
    add_three_col_table(
        doc,
        ["Usługa", "Odpowiedzialność i własne dane", "Przyszła komunikacja"],
        [
            ["Pierwsza — implementowana", "Wpisz zakres i dane, za które odpowiada.", "Jaką informację może udostępniać przez API?"],
            ["Druga — planowana", "Wpisz odrębną odpowiedzialność i jej dane.", "Jak pierwsza usługa mogłaby się z nią skontaktować?"],
            ["Opcjonalna trzecia", "Dodaj tylko wtedy, gdy wynika z problemu.", "Jaki byłby jej odbiorca lub nadawca komunikatu?"],
        ],
    )
    add_spacer(doc)
    add_two_col_table(doc, "Właściciel danych", "Wybierz jeden ważny rodzaj danych i napisz, która usługa może go tworzyć oraz zmieniać.", lines=3)
    add_spacer(doc)
    add_two_col_table(doc, "Przyszły kontrakt", "Naszkicuj jedno wywołanie HTTP między usługami: nadawca, odbiorca, metoda, ścieżka, przykładowe żądanie i odpowiedź.", lines=5)
    add_spacer(doc)
    add_two_col_table(doc, "Niedostępność odbiorcy", "Co powinien zrobić nadawca, gdy druga usługa nie odpowiada? Zapisz założenie, nie implementuj jeszcze komunikacji.", lines=3)

    add_page_break(doc)
    add_heading(doc, "4  Pierwsza usługa i jej API")
    add_instruction(doc, "Zaprojektuj małą usługę, którą da się uruchomić niezależnie. Wybierz jedną regułę biznesową dotyczącą jej własnych danych.")
    add_spacer(doc)
    add_two_col_table(doc, "Główny obiekt domenowy", "Podaj nazwę obiektu i jego najważniejsze pola. Nie twórz jeszcze całej bazy danych.", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Reguła biznesowa", "Zapisz regułę wykraczającą poza wymagane pola, typy i unikalność identyfikatora.", lines=3)
    add_spacer(doc)
    add_two_col_table(doc, "Dozwolone i niedozwolone", "Podaj przykład operacji dozwolonej, niedozwolonej oraz oczekiwany kod błędu.", lines=3)
    add_spacer(doc)
    add_three_col_table(
        doc,
        ["Metoda", "Endpoint", "Znaczenie i odpowiedź"],
        [
            ["GET", "/...", "Lista zasobów; sukces 200."],
            ["GET", "/.../{id}", "Jeden zasób; brak zasobu 404."],
            ["POST albo PATCH", "/...", "Zapis; sukces 201 albo 200, konflikt 409, błędne dane 422."],
            ["GET", "/health", "Sprawdzenie działania usługi; 200."],
        ],
    )
    add_spacer(doc)
    add_two_col_table(doc, "Przykładowe żądanie i odpowiedź", "Wpisz jeden przykładowy JSON żądania i odpowiedzi sukcesu. Dodaj odpowiedź błędu dla swojej reguły.", lines=5)
    add_spacer(doc)
    add_two_col_table(doc, "Awaria usługi", "Co przestanie działać, gdy proces tej usługi zostanie zatrzymany? Co nadal mogłoby działać w pozostałej części systemu?", lines=3)

    add_heading(doc, "5  MVP i kryteria sukcesu")
    add_spacer(doc)
    add_two_col_table(doc, "Minimalna wersja projektu", "Wpisz najmniejszy zakres, który można uruchomić i pokazać. Wymień funkcje, które świadomie odkładasz.", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Kryteria sukcesu", "Wpisz 2–4 sprawdzalne kryteria poprawności i użyteczności rozwiązania.", lines=4)
    add_spacer(doc)
    add_two_col_table(doc, "Czy mikroserwisy mają sens", "Porównaj krótko ten pomysł z monolitem. Wskaż jedną korzyść i jeden koszt podziału na usługi.", lines=4)

    add_page_break(doc)
    add_heading(doc, "6  Zespół i odpowiedzialność")
    add_instruction(doc, "Każda osoba powinna mieć własny, mierzalny rezultat. Wpisz konkretny artefakt, a nie ogólne „pomoc przy projekcie”.")
    add_spacer(doc)
    add_three_col_table(
        doc,
        ["Osoba", "Zakres odpowiedzialności", "Mierzalny rezultat i sposób sprawdzenia"],
        [
            ["1.", "Wpisz obszar odpowiedzialności.", "Np. usługa X, testy Y; jak prowadzący to sprawdzi?"],
            ["2.", "Wpisz obszar odpowiedzialności.", "Np. usługa X, dokumentacja Y; jak prowadzący to sprawdzi?"],
            ["3.", "Opcjonalnie.", "Wpisz rezultat albo pozostaw pustą komórkę."],
            ["4.", "Opcjonalnie.", "Wpisz rezultat albo pozostaw pustą komórkę."],
        ],
    )

    add_heading(doc, "7  Ryzyka i wykorzystanie narzędzi")
    add_spacer(doc)
    add_two_col_table(doc, "Największe ryzyka", "Wypisz do trzech ryzyk, ich możliwy skutek oraz pomysł na ograniczenie ryzyka. Uwzględnij dane, zakres i integrację.", lines=5)
    add_spacer(doc)
    add_three_col_table(
        doc,
        ["Narzędzie", "Do czego użyto", "Co zweryfikowano"],
        [
            ["AI / dokumentacja / inne", "Wpisz, do czego wykorzystano narzędzie.", "Jak sprawdzono poprawność informacji lub kodu?"],
            ["", "", ""],
            ["", "", ""],
        ],
    )
    add_spacer(doc)
    add_instruction(doc, "Przed kolejnymi zajęciami sprawdźcie: czy problem jest konkretny, odbiorca jasno określony, granica pierwszej usługi uzasadniona, MVP wykonalne, a kryteria sukcesu i ryzyka sprawdzalne.")

    footer = doc.sections[0].footer
    if footer.paragraphs:
        p = footer.paragraphs[0]
        fill_paragraph(p, "Projektowanie aplikacji webowych opartych o mikrousługi  |  LAB1  |  Karta pomysłu", "Footer", color=BLUE, size=8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))


if __name__ == "__main__":
    build()
